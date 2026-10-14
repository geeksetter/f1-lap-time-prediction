"""Render the editable practice report from Markdown and local figures.

Optional dependency: python-docx==1.2.0. Export PDF with Word/LibreOffice.
The numerical results are read from the verified report source, not recalculated.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def text_value(text):
    text = re.sub(r"!?\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


def font(style, name, size):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def table(document, rows):
    parsed = [row.strip().strip("|").split("|") for row in rows]
    parsed = [row for row in parsed if not all(re.fullmatch(r"[\s:\-]+", x) for x in row)]
    count = max(map(len, parsed))
    result = document.add_table(rows=0, cols=count)
    result.style = "Table Grid"
    for i, row in enumerate(parsed):
        cells = result.add_row().cells
        no_split = OxmlElement("w:cantSplit")
        result.rows[-1]._tr.get_or_add_trPr().append(no_split)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            result.rows[-1]._tr.get_or_add_trPr().append(repeat)
        for j, value in enumerate(row):
            paragraph = cells[j].paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(3)
            if len(parsed) <= 5 and i < len(parsed) - 1:
                paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run(text_value(value.strip()))
            run.font.name = "Times New Roman"
            run.font.size = Pt(9 if count > 6 else 10)
            run.bold = i == 0
    document.add_paragraph()


def build():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    normal = document.styles["Normal"]
    font(normal, "Times New Roman", 14)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 18), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 14)]:
        font(document.styles[name], "Times New Roman", size)
        document.styles[name].paragraph_format.keep_with_next = True
    for style in document.styles:
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = (ROOT / "docs/practice_report.md").read_text().splitlines()
    i, image_number, title_seen, diagram_seen = 0, 0, False, False
    in_contents = False
    while i < len(lines):
        line = lines[i]
        if line == "<!-- pagebreak -->":
            pass  # Top-level headings already start a page; avoid duplicate breaks.
        elif line.startswith("```"):
            kind = line[3:]
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            if kind == "mermaid":
                if not diagram_seen:
                    p = document.add_paragraph()
                    p.paragraph_format.keep_with_next = True
                    p.add_run().add_picture(str(ROOT / "docs/figures/er_schema.png"), width=Cm(16.3))
                    document.add_paragraph("Рисунок 1 Логическая схема основных связей по SQL DDL. Это иллюстрация схемы, а не скриншот DBeaver.")
                    diagram_seen = True
                    image_number += 1
            else:
                for source in block:
                    p = document.add_paragraph()
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.widow_control = False
                    run = p.add_run(source or " ")
                    run.font.name = "Courier New"
                    run.font.size = Pt(8.5)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            table(document, rows)
            continue
        elif re.match(r"!\[.*\]\(.*\)", line):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            image_number += 1
            p = document.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(ROOT / "docs" / match[2]), width=Cm(16.3))
            document.add_paragraph(f"Рисунок {image_number} {match[1]}")
        elif line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 3)
            text = text_value(line.lstrip("# "))
            text = text.replace(":", " ").replace("—", " ")
            if not title_seen:
                p = document.add_paragraph(text, "Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_seen = True
            else:
                p = document.add_heading(text, level)
                if level == 1:
                    p.paragraph_format.page_break_before = True
                    in_contents = text == "Содержание"
        elif line.strip():
            line = re.sub(r"^[-*] ", "• ", line)
            p = document.add_paragraph(text_value(line))
            if in_contents:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(5)
                for run in p.runs:
                    run.font.size = Pt(12)
        i += 1
    document.core_properties.title = "Учебная практика Прогноз времени круга Формулы-1"
    document.core_properties.author = "geeksetter"
    document.core_properties.subject = "Реальные OpenF1 данные и собственная ООП-регрессия через PostgreSQL"
    target = ROOT / "docs/practice_report.docx"
    document.save(target)
    print(f"Saved {target.name}")


if __name__ == "__main__":
    build()
